#!/usr/bin/env python3
"""
사내 작업안전가이드 등록 도구
PDF/Word/Excel 파일을 텍스트로 변환하여 .claude/safety-guide/guide.txt 에 저장합니다.

사용법:
  python3 register_guide.py <파일경로>
  python3 register_guide.py /path/to/안전가이드.pdf
  python3 register_guide.py /path/to/안전가이드.docx
  python3 register_guide.py /path/to/안전가이드.xlsx

여러 파일을 순서대로 등록하면 내용이 누적됩니다:
  python3 register_guide.py 1장_일반안전.pdf
  python3 register_guide.py 2장_기계안전.pdf
  python3 register_guide.py 3장_화학안전.pdf
"""

import sys
import os

OUTPUT_DIR  = os.path.join(os.path.dirname(__file__))
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "guide.txt")


def extract_pdf(path):
    try:
        import pdfplumber
        texts = []
        with pdfplumber.open(path) as pdf:
            total = len(pdf.pages)
            print(f"  PDF 페이지 수: {total}")
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    texts.append(f"[PAGE {i}]\n{text}")
                if i % 10 == 0:
                    print(f"  처리 중... {i}/{total} 페이지")
        return "\n\n".join(texts)
    except Exception as e:
        print(f"  pdfplumber 실패: {e}, pypdf로 재시도합니다.")
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            texts = []
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    texts.append(f"[PAGE {i}]\n{text}")
            return "\n\n".join(texts)
        except Exception as e2:
            raise RuntimeError(f"PDF 변환 실패: {e2}")


def extract_docx(path):
    try:
        import docx
    except ImportError:
        os.system("pip3 install python-docx -q")
        import docx
    doc = docx.Document(path)
    texts = []
    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                texts.append(row_text)
    return "\n".join(texts)


def extract_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        os.system("pip3 install openpyxl -q")
        import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    texts = []
    for sheet in wb.worksheets:
        texts.append(f"[시트: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                texts.append(row_text)
    return "\n".join(texts)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    filepath = sys.argv[1]
    if not os.path.exists(filepath):
        print(f"❌ 파일을 찾을 수 없습니다: {filepath}")
        sys.exit(1)

    ext = os.path.splitext(filepath)[1].lower()
    filename = os.path.basename(filepath)
    print(f"📄 파일 처리 중: {filename}")

    if ext == ".pdf":
        text = extract_pdf(filepath)
    elif ext in (".docx", ".doc"):
        text = extract_docx(filepath)
    elif ext in (".xlsx", ".xls"):
        text = extract_xlsx(filepath)
    elif ext == ".txt":
        with open(filepath, encoding="utf-8", errors="ignore") as f:
            text = f.read()
    else:
        print(f"❌ 지원하지 않는 파일 형식입니다: {ext}")
        print("   지원 형식: PDF, DOCX, XLSX, TXT")
        sys.exit(1)

    # 누적 저장 (파일 구분자 포함)
    mode = "a" if os.path.exists(OUTPUT_FILE) else "w"
    with open(OUTPUT_FILE, mode, encoding="utf-8") as f:
        f.write(f"\n{'='*60}\n")
        f.write(f"출처: {filename}\n")
        f.write(f"{'='*60}\n\n")
        f.write(text)
        f.write("\n")

    char_count = len(text)
    print(f"✅ 등록 완료: {OUTPUT_FILE}")
    print(f"   추출 글자 수: {char_count:,}자")
    print(f"   저장 모드: {'추가' if mode == 'a' else '신규'}")

    if os.path.exists(OUTPUT_FILE):
        total_size = os.path.getsize(OUTPUT_FILE)
        print(f"   guide.txt 누적 크기: {total_size:,} bytes")


if __name__ == "__main__":
    main()
